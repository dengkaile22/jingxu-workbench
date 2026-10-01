"""Stable account identity and duplicate-safe crawl exports, without network access."""

from __future__ import annotations

import json
from dataclasses import replace

from core.analyzer import ProductAnalyzer
from core.exporter import DataExporter, _build_account_rows
from core.models import (
    AccountTarget,
    OccurrenceFilterConfig,
    TimeFilterConfig,
    VideoProductItem,
)
from openpyxl import load_workbook


def _account(
    url: str = "https://v.douyin.com/first/",
    *,
    platform: str = "douyin",
    account_id: str = "MS4wLjABAAAA_account1",
    name: str = "池恩陈.",
    status: str = "已完成",
) -> AccountTarget:
    return AccountTarget(
        platform=platform,
        url_or_id=url,
        name=name,
        sec_uid=account_id if platform == "douyin" else "",
        user_id=account_id if platform == "kuaishou" else "",
        status=status,
    )


def _item(
    video_id: str = "video-1",
    *,
    platform: str = "douyin",
    account_id: str = "MS4wLjABAAAA_account1",
    name: str = "池恩陈.",
    product_id: str = "product-1",
    product_name: str = "生姜防脱洗发水500ml",
) -> VideoProductItem:
    return VideoProductItem(
        platform=platform,
        account_name=name,
        account_id=account_id,
        video_id=video_id,
        video_title=f"商品介绍 {video_id}",
        video_url=f"https://www.{platform}.com/video/{video_id}",
        publish_time=1790683200,
        publish_time_str="2026-09-29 20:00:00",
        product_id=product_id,
        product_name=product_name,
    )


def _aggregate(items: list[VideoProductItem]):
    analyzer = ProductAnalyzer(
        OccurrenceFilterConfig(min_account_count=1, enable_fuzzy_cluster=False)
    )
    return analyzer.aggregate_and_filter(items)[1]


def test_stable_account_key_uses_platform_and_sec_uid() -> None:
    from core.record_identity import stable_account_key

    first = _account()
    second = _account("https://v.douyin.com/second/")
    assert stable_account_key(first) == stable_account_key(second)
    platform, identity = stable_account_key(first)
    assert platform == "douyin"
    assert first.sec_uid in identity
    assert stable_account_key(_account(platform="kuaishou")) != stable_account_key(first)


def test_unknown_accounts_do_not_merge_by_shared_nickname() -> None:
    from core.record_identity import deduplicate_accounts, stable_account_key

    accounts = [
        _account("https://v.douyin.com/first/", account_id=""),
        _account("https://v.douyin.com/second/", account_id=""),
    ]
    assert stable_account_key(accounts[0]) != stable_account_key(accounts[1])
    assert len(deduplicate_accounts(accounts)) == 2


def test_two_short_links_with_same_sec_uid_export_one_statistics_row() -> None:
    accounts = [_account(), _account("https://v.douyin.com/second/")]
    rows = _build_account_rows(accounts, [_item()])

    assert len(rows) == 1
    assert rows[0][4:] == [1, 1]


def test_twenty_five_duplicate_video_records_count_once_in_statistics() -> None:
    items = [_item(f"video-{index}") for index in range(25)]
    duplicates = [replace(item) for item in items]
    rows = _build_account_rows([_account()], items + duplicates)

    assert len(rows) == 1
    assert rows[0][4:] == [25, 25]


def test_excel_detail_count_matches_deduplicated_rows(tmp_path) -> None:
    items = [_item(f"video-{index}") for index in range(25)]
    duplicates = [replace(item) for item in items]
    destination = tmp_path / "results.xlsx"
    DataExporter.export_to_excel(
        str(destination),
        [],
        items + duplicates,
        [_account(), _account("https://v.douyin.com/second/")],
    )

    workbook = load_workbook(destination, read_only=True)
    try:
        detail = workbook["全部带货商品明细(25条)"]
        assert detail.max_row == 26
        statistics = workbook["对标账号爬取统计"]
        assert statistics.max_row == 2
        assert list(statistics.values)[1][4:] == (25, 25)
    finally:
        workbook.close()


def test_same_nickname_distinct_ids_are_two_accounts() -> None:
    products = _aggregate([
        _item("video-1", account_id="account-a"),
        _item("video-2", account_id="account-b"),
    ])

    assert len(products) == 1
    assert products[0].account_count == 2
    assert products[0].video_count == 2


def test_same_nickname_different_ids_do_not_share_statistics() -> None:
    accounts = [_account(account_id="account-a"), _account(account_id="account-b")]
    items = [
        _item("a-1", account_id="account-a"),
        _item("a-2", account_id="account-a"),
        _item("b-1", account_id="account-b"),
    ]
    rows = _build_account_rows(accounts, items)

    assert len(rows) == 2
    assert [row[4:] for row in rows] == [[2, 2], [1, 1]]


def test_renamed_account_still_counts_as_one_account() -> None:
    products = _aggregate([
        _item("video-1", name="原昵称"),
        _item("video-2", name="新昵称"),
    ])

    assert len(products) == 1
    assert products[0].account_count == 1
    assert products[0].video_count == 2


def test_duplicated_items_are_removed_before_product_aggregation() -> None:
    items = [_item(f"video-{index}") for index in range(25)]
    products = _aggregate(items + [replace(item) for item in items])

    assert len(products) == 1
    assert products[0].account_count == 1
    assert products[0].video_count == 25
    assert len(products[0].items) == 25


def test_same_nickname_across_platforms_has_separate_statistics() -> None:
    accounts = [_account(), _account(platform="kuaishou")]
    items = [
        _item("d-1"),
        _item("d-2"),
        _item("k-1", platform="kuaishou"),
    ]
    rows = _build_account_rows(accounts, items)

    assert [row[0] for row in rows] == ["douyin", "kuaishou"]
    assert [row[4:] for row in rows] == [[2, 2], [1, 1]]


def test_unresolved_ids_do_not_match_all_items_with_empty_ids_or_same_name() -> None:
    accounts = [
        _account("https://v.douyin.com/first/", account_id=""),
        _account("https://v.douyin.com/second/", account_id=""),
    ]
    rows = _build_account_rows(accounts, [_item(account_id="")])

    assert len(rows) == 2
    assert [row[4:] for row in rows] == [[0, 0], [0, 0]]


def test_item_deduplication_preserves_multiple_products_in_one_video() -> None:
    from core.record_identity import deduplicate_items

    first = _item(product_id="product-a")
    second = _item(product_id="product-b", product_name="氨基酸护发素500ml")
    items = deduplicate_items([first, second, replace(first), replace(second)])

    assert len(items) == 2
    assert {item.product_id for item in items} == {"product-a", "product-b"}
    assert _build_account_rows([_account()], items)[0][4:] == [1, 2]


def test_item_deduplication_preserves_distinct_products_without_product_ids() -> None:
    from core.record_identity import deduplicate_items

    first = _item(product_id="", product_name="生姜洗发水500ml")
    second = _item(product_id="", product_name="氨基酸护发素500ml")
    items = deduplicate_items([first, second, replace(first), replace(second)])

    assert len(items) == 2


def test_item_deduplication_is_scoped_by_platform() -> None:
    from core.record_identity import deduplicate_items

    douyin = _item()
    kuaishou = _item(platform="kuaishou")
    items = deduplicate_items([douyin, kuaishou, replace(douyin), replace(kuaishou)])

    assert len(items) == 2
    assert {item.platform for item in items} == {"douyin", "kuaishou"}


def test_json_preserves_account_identity_and_reports_partial_result(tmp_path) -> None:
    destination = tmp_path / "partial.json"
    account = _account(status="部分完成")
    DataExporter.export_to_json(
        str(destination), [], [_item(), _item()], accounts=[account]
    )
    data = json.loads(destination.read_text(encoding="utf-8"))

    assert data["total_items_count"] == 1
    assert len(data["all_items"]) == 1
    assert len(data["accounts"]) == 1
    assert data["accounts"][0]["sec_uid"] == account.sec_uid
    assert data["accounts"][0]["status"] == "部分完成"
    assert data["crawl_complete"] is False
    assert data["message"]


def test_json_reports_completion_for_all_completed_accounts(tmp_path) -> None:
    destination = tmp_path / "complete.json"
    DataExporter.export_to_json(str(destination), [], [_item()], accounts=[_account()])
    data = json.loads(destination.read_text(encoding="utf-8"))

    assert data["accounts"][0]["sec_uid"] == _account().sec_uid
    assert data["crawl_complete"] is True


def test_kuaishou_query_accounts_remain_distinct():
    from core.record_identity import deduplicate_accounts, stable_account_key
    accounts = [AccountTarget("kuaishou", f"https://www.kuaishou.com/?userId={identity}") for identity in ("account-a", "account-b")]
    assert len(deduplicate_accounts(accounts)) == 2
    assert stable_account_key(accounts[0]) == stable_account_key(AccountTarget("kuaishou", "https://www.kuaishou.com/profile/account-a"))


def test_invalid_account_url_can_be_exported_as_a_failed_attempt(tmp_path):
    target = AccountTarget("douyin", "https://[invalid-url", status="解析失败")
    destination = tmp_path / "failed.json"
    DataExporter.export_to_json(str(destination), [], [], accounts=[target])
    assert json.loads(destination.read_text(encoding="utf-8"))["crawl_complete"] is False


def test_cached_kuaishou_identity_does_not_resolve_expired_share_link(monkeypatch):
    from types import SimpleNamespace

    from core import kuaishou_crawler

    def unexpected_resolution(_target):
        raise AssertionError("cached user ID must take priority")
    monkeypatch.setattr(kuaishou_crawler, "resolve_kuaishou_url_or_userid", unexpected_resolution)
    target = AccountTarget("kuaishou", "https://v.kuaishou.com/expired/", user_id="cached-user")
    page = SimpleNamespace(
        listen=SimpleNamespace(start=lambda _targets: None), get=lambda *_args, **_kwargs: None,
        ele=lambda *_args, **_kwargs: None,
    )
    kuaishou_crawler.KuaishouCrawler(SimpleNamespace(get_page=lambda: page)).crawl_account(
        target, TimeFilterConfig(mode="all"), max_scrolls=0,
    )
    assert target.user_id == "cached-user"
