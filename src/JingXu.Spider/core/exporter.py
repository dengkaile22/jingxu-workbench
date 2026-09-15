"""
Data exporter for crawl and analysis results (Excel, CSV, JSON).
"""
import os
import csv
import json
from typing import List, Dict, Any, Optional
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core.models import RetainedProduct, VideoProductItem, AccountTarget

RETAINED_HEADERS = [
    "序号", "商品名称", "清洗归一名称", "涉及平台", "共同带货账号数",
    "带货视频总数", "带货账号列表", "参考价格区间", "示例视频链接",
]

DETAIL_HEADERS = [
    "序号", "平台", "账号名称", "账号ID", "视频ID", "视频标题", "发布时间",
    "商品原始名称", "商品清洗名称", "商品价格", "视频链接", "商品链接",
]

ACCOUNT_HEADERS = [
    "平台", "账号名称", "账号主页/ID", "爬取状态", "发现带货视频数", "发现商品记录数",
]


def _build_retained_rows(retained_products: List[RetainedProduct]) -> List[List[Any]]:
    rows = []
    for idx, p in enumerate(retained_products, 1):
        rows.append([
            idx,
            p.product_name,
            p.cleaned_name,
            p.platform,
            p.account_count,
            p.video_count,
            ", ".join(p.accounts),
            p.price_range,
            p.sample_video_url,
        ])
    return rows


def _build_detail_rows(all_items: List[VideoProductItem]) -> List[List[Any]]:
    rows = []
    for idx, it in enumerate(all_items, 1):
        rows.append([
            idx,
            it.platform,
            it.account_name,
            it.account_id,
            it.video_id,
            it.video_title,
            it.publish_time_str,
            it.product_name,
            it.cleaned_name,
            it.price,
            it.video_url,
            it.detail_url,
        ])
    return rows


def _build_account_rows(
    accounts: List[AccountTarget],
    all_items: List[VideoProductItem],
) -> List[List[Any]]:
    rows = []
    for acc in accounts:
        acc_items = [
            it for it in all_items
            if it.account_name == acc.name or it.account_id in (acc.sec_uid, acc.user_id, acc.url_or_id)
        ]
        acc_videos = set(it.video_id for it in acc_items)
        rows.append([
            acc.platform,
            acc.name or acc.url_or_id,
            acc.url_or_id,
            acc.status,
            len(acc_videos),
            len(acc_items),
        ])
    return rows


def _write_sheet(worksheet, headers: List[str], rows: List[List[Any]]) -> None:
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Microsoft YaHei", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    worksheet.views.sheetView[0].showGridLines = True

    for col_idx in range(1, len(headers) + 1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in worksheet[col_letter]:
            if cell.row > 1:
                cell.font = cell_font
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")
            val_str = str(cell.value or "")
            char_len = sum(2 if ord(c) > 127 else 1 for c in val_str)
            if char_len > max_len:
                max_len = char_len
        worksheet.column_dimensions[col_letter].width = max(10, min(max_len + 4, 60))

    worksheet.row_dimensions[1].height = 26
    for r in range(2, worksheet.max_row + 1):
        worksheet.row_dimensions[r].height = 20


class DataExporter:
    """Exports crawled data and retained product analysis into Excel, CSV, or JSON."""

    @staticmethod
    def export_to_excel(
        file_path: str,
        retained_products: List[RetainedProduct],
        all_items: List[VideoProductItem],
        accounts: List[AccountTarget] = None
    ) -> str:
        """Export analysis results to a styled Excel file with multiple sheets."""
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)

        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = f"保留相同商品汇总({len(retained_products)}款)"
        _write_sheet(summary_sheet, RETAINED_HEADERS, _build_retained_rows(retained_products))

        detail_sheet = workbook.create_sheet(f"全部带货商品明细({len(all_items)}条)")
        _write_sheet(detail_sheet, DETAIL_HEADERS, _build_detail_rows(all_items))

        account_rows = _build_account_rows(accounts or [], all_items)
        if account_rows:
            account_sheet = workbook.create_sheet("对标账号爬取统计")
            _write_sheet(account_sheet, ACCOUNT_HEADERS, account_rows)

        workbook.save(file_path)
        return file_path

    @staticmethod
    def export_to_csv(file_path: str, retained_products: List[RetainedProduct]) -> str:
        """Export retained products to CSV."""
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        with open(file_path, "w", encoding="utf_8_sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=RETAINED_HEADERS, extrasaction="ignore")
            writer.writeheader()
            for idx, p in enumerate(retained_products, 1):
                writer.writerow({
                    "序号": idx,
                    "商品名称": p.product_name,
                    "清洗归一名称": p.cleaned_name,
                    "涉及平台": p.platform,
                    "共同带货账号数": p.account_count,
                    "带货视频总数": p.video_count,
                    "带货账号列表": ", ".join(p.accounts),
                    "参考价格区间": p.price_range,
                    "示例视频链接": p.sample_video_url,
                })
        return file_path

    @staticmethod
    def export_to_json(
        file_path: str,
        retained_products: List[RetainedProduct],
        all_items: List[VideoProductItem],
        excel_path: Optional[str] = None,
    ) -> str:
        """Export full data to JSON."""
        data = {
            "exported_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "retained_count": len(retained_products),
            "total_items_count": len(all_items),
            "excel_path": excel_path or "",
            "retained_products": [p.to_summary_dict() for p in retained_products],
            "all_items": [it.to_dict() for it in all_items],
        }
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        with open(file_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        return file_path
