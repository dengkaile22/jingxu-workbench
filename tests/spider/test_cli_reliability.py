"""Offline CLI integration: response diagnostics, completion and stable identities."""

from __future__ import annotations

import importlib
import json
from dataclasses import replace
from types import SimpleNamespace

from core.diagnostics import PreflightReport
from core.extraction import ExtractionDiagnostic
from core.models import (
    AccountTarget,
    OccurrenceFilterConfig,
    TimeFilterConfig,
    VideoProductItem,
)
from openpyxl import load_workbook
from DrissionPage.errors import PageDisconnectedError

SEC_UID = "MS4wLjABAAAA_stable-account"


def _item(index: int) -> VideoProductItem:
    return VideoProductItem(
        platform="douyin",
        account_name="池恩陈.",
        account_id=SEC_UID,
        video_id=f"video-{index}",
        video_title=f"商品介绍 {index}",
        video_url=f"https://www.douyin.com/video/video-{index}",
        publish_time=1790683200,
        publish_time_str="2026-09-29 20:00:00",
        product_id="product-1",
        product_name="生姜防脱洗发水500ml",
    )


def _run_offline_cli(
    monkeypatch,
    tmp_path,
    *,
    input_urls=None,
    status="已完成",
    response_issues=None,
    video_count=1,
    fail_after_checkpoint=False,
    cancel_after_checkpoint=False,
):
    """Run real CLI analysis and exporters while replacing external browser work."""
    cli = importlib.import_module("cli")
    accounts = [
        AccountTarget(platform="douyin", url_or_id=url, name="池恩陈.")
        for url in (input_urls or ["https://v.douyin.com/first/"])
    ]
    config = SimpleNamespace(
        accounts=accounts,
        time_config=TimeFilterConfig(mode="all"),
        filter_config=OccurrenceFilterConfig(
            min_account_count=1, enable_fuzzy_cluster=False
        ),
        headless=True,
        max_scrolls_per_account=25,
        export_dir=str(tmp_path),
    )
    browsers = []
    crawl_calls = []

    class OfflineBrowserManager:
        def __init__(self, **kwargs):
            self.closed = False
            browsers.append(self)

        def close(self):
            self.closed = True

        def stop(self):
            self.close()

    class OfflineDouyinCrawler:
        def __init__(self, *, browser_manager, log_callback):
            self.browser_manager = browser_manager
            self.diagnostics = []
            self.response_issues = list(response_issues or [])

        def crawl_account(self, target, time_config, max_scrolls):
            crawl_calls.append(target.url_or_id)
            target.sec_uid = SEC_UID
            target.status = status
            target.scraped_videos = video_count
            target.found_products = video_count
            items = [_item(index) for index in range(video_count)]
            self.diagnostics.extend(
                ExtractionDiagnostic(
                    platform="douyin", video_id=item.video_id, extracted=1,
                    known_paths_hit=["anchor_info.extra"], scanned_nodes=1,
                )
                for item in items
            )
            if len(crawl_calls) == 1 and (fail_after_checkpoint or cancel_after_checkpoint):
                items[0].video_id = 'checkpoint-before-disconnect'
                self.on_checkpoint(items)
                if cancel_after_checkpoint:
                    (tmp_path / 'cancel.request').write_text('cancel', encoding='utf-8')
                raise PageDisconnectedError('synthetic-private-token')
            # Deliberately return repeats, including across different input links:
            # the CLI/export path must remain duplicate-safe independently.
            return items + [replace(item) for item in items]

    class OfflineKuaishouCrawler:
        def __init__(self, *, browser_manager, log_callback):
            self.browser_manager = browser_manager
            self.diagnostics = []

        def crawl_account(self, *args):
            raise AssertionError("Unexpected Kuaishou crawl for Douyin-only input")

    destination = tmp_path / "results.json"
    workbook_path = tmp_path / "results.xlsx"
    monkeypatch.setattr(cli, "AppConfig", lambda path: config)
    monkeypatch.setattr(cli, "BrowserManager", OfflineBrowserManager)
    monkeypatch.setattr(cli, "DouyinCrawler", OfflineDouyinCrawler)
    monkeypatch.setattr(cli, "KuaishouCrawler", OfflineKuaishouCrawler)
    monkeypatch.setattr(cli, "run_preflight", lambda **kwargs: PreflightReport())
    monkeypatch.setattr(cli, "describe_login_state", lambda *args: "登录态正常")
    monkeypatch.setattr(cli, "load_login_state", lambda *args: {"douyin": {"status": "ok"}})
    monkeypatch.setattr(cli, "record_login_state", lambda *args, **kwargs: {})
    monkeypatch.setattr(cli, "_probe_login_state", lambda crawler: "ok")
    monkeypatch.setattr(
        cli.sys, "argv",
        ["cli.py", "--config", str(tmp_path / "config.json"),
         "--export", str(workbook_path), "--json-out", str(destination),
         "--profile-dir", str(tmp_path / "browser-profile"),
         *(["--cancel-file", str(tmp_path / 'cancel.request')] if cancel_after_checkpoint else [])],
    )

    cli.run_cli()
    assert len(browsers) == 1
    assert browsers[0].closed is True
    return json.loads(destination.read_text(encoding="utf-8")), workbook_path, crawl_calls


def test_cli_carries_response_failures_into_json_diagnostics(monkeypatch, tmp_path) -> None:
    response_issue = {
        "level": "warning",
        "code": "douyin_response_non_json",
        "message": "抖音视频响应为验证页面，未完整抓取",
        "suggestion": "完成页面验证后重试",
    }
    data, _, _ = _run_offline_cli(
        monkeypatch, tmp_path, status="部分完成", response_issues=[response_issue]
    )

    issues = data["diagnostics"]["issues"]
    assert response_issue in issues
    assert any(issue["code"] == "crawl_incomplete" for issue in issues)


def test_cli_partial_result_keeps_identity_and_incomplete_message(monkeypatch, tmp_path) -> None:
    data, _, _ = _run_offline_cli(monkeypatch, tmp_path, status="部分完成")

    assert data["crawl_complete"] is False
    assert data["message"].startswith("抓取未完整结束")
    assert not data["message"].startswith("抓取完成")
    assert data["total_items_count"] == 1
    assert data["accounts"][0]["sec_uid"] == SEC_UID
    assert data["accounts"][0]["status"] == "部分完成"
    assert data["account_resolutions"][0]["sec_uid"] == SEC_UID
    assert data["account_resolutions"][0]["url_or_id"] == "https://v.douyin.com/first/"


def test_cli_two_resolved_links_and_duplicate_records_export_once(monkeypatch, tmp_path) -> None:
    urls = ["https://v.douyin.com/first/", "https://v.douyin.com/second/"]
    data, workbook_path, calls = _run_offline_cli(
        monkeypatch, tmp_path, input_urls=urls, video_count=25
    )

    assert calls == urls
    assert data["crawl_complete"] is True
    assert data["total_items_count"] == 25
    assert len(data["all_items"]) == 25
    assert len(data["accounts"]) == 1
    assert data["accounts"][0]["sec_uid"] == SEC_UID
    assert len(data["account_resolutions"]) == 2
    assert {account["url_or_id"] for account in data["account_resolutions"]} == set(urls)
    assert {account["sec_uid"] for account in data["account_resolutions"]} == {SEC_UID}
    assert data["retained_count"] == 1
    assert data["retained_products"][0]["带货账号数"] == 1
    assert data["retained_products"][0]["带货视频数"] == 25

    workbook = load_workbook(workbook_path, read_only=True)
    try:
        assert workbook["全部带货商品明细(25条)"].max_row == 26
        statistics = workbook["对标账号爬取统计"]
        assert statistics.max_row == 2
        assert list(statistics.values)[1][4:] == (25, 25)
    finally:
        workbook.close()


def test_cli_disconnect_preserves_checkpoint_and_continues_next_account(monkeypatch, tmp_path, capsys):
    urls = ['https://v.douyin.com/first/', 'https://v.douyin.com/second/']
    data, workbook_path, calls = _run_offline_cli(
        monkeypatch, tmp_path, input_urls=urls, fail_after_checkpoint=True
    )
    assert calls == urls
    assert data['total_items_count'] == 2
    assert {item['视频ID'] for item in data['all_items']} == {'checkpoint-before-disconnect', 'video-0'}
    assert data['crawl_complete'] is False
    assert data['account_resolutions'][0]['status'] == '部分完成'
    assert data['account_resolutions'][1]['status'] == '已完成'
    assert any(issue['code'] == 'page_disconnected' for issue in data['diagnostics']['issues'])
    assert workbook_path.is_file()
    output = capsys.readouterr().out
    assert 'Traceback' not in output and 'synthetic-private-token' not in output


def test_cli_cancel_disconnect_exports_partial_without_continuing(monkeypatch, tmp_path, capsys):
    urls = ['https://v.douyin.com/first/', 'https://v.douyin.com/second/']
    data, workbook_path, calls = _run_offline_cli(
        monkeypatch, tmp_path, input_urls=urls, cancel_after_checkpoint=True
    )
    assert calls == urls[:1]
    assert data['total_items_count'] == 1 and data['crawl_complete'] is False
    assert data['account_resolutions'][0]['status'] == '已停止'
    assert workbook_path.is_file()
    assert 'Traceback' not in capsys.readouterr().out
