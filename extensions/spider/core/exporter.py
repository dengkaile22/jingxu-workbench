"""
Data exporter for crawl and analysis results (Excel, CSV, JSON).
"""

import os
import csv
import json
import tempfile
from typing import List, Dict, Any, Optional
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter

from core.models import RetainedProduct, VideoProductItem, AccountTarget
from core.record_identity import (
    account_profile_url,
    deduplicate_accounts,
    deduplicate_items,
    stable_id_for,
)

RETAINED_HEADERS = [
    "序号",
    "商品名称",
    "清洗归一名称",
    "涉及平台",
    "共同带货账号数",
    "带货视频总数",
    "示例视频点赞量",
    "带货账号列表",
    "参考价格区间",
    "示例视频日期",
    "示例视频链接",
]

DETAIL_HEADERS = [
    "序号",
    "平台",
    "账号名称",
    "账号ID",
    "视频ID",
    "视频标题",
    "发布时间",
    "商品原始名称",
    "商品清洗名称",
    "商品价格",
    "视频链接",
    "商品链接",
]

ACCOUNT_HEADERS = [
    "平台",
    "账号名称",
    "账号主页/ID",
    "爬取状态",
    "发现带货视频数",
    "发现商品记录数",
]


def _build_retained_rows(retained_products: List[RetainedProduct]) -> List[List[Any]]:
    rows = []
    for idx, p in enumerate(retained_products, 1):
        rows.append(
            [
                idx,
                p.product_name,
                p.cleaned_name,
                p.platform,
                p.account_count,
                p.video_count,
                p.sample_video_like_count,
                ", ".join(p.accounts),
                p.price_range,
                _excel_date(p.sample_video_date),
                p.sample_video_url,
            ]
        )
    return rows


def _excel_date(value: str) -> Any:
    """Write valid YYYY-MM-DD values as native Excel dates for date filtering."""
    if not value:
        return ""
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError):
        return value


def _build_detail_rows(all_items: List[VideoProductItem]) -> List[List[Any]]:
    rows = []
    for idx, it in enumerate(all_items, 1):
        rows.append(
            [
                idx,
                it.platform,
                it.account_name,
                it.account_id,
                it.video_id,
                it.video_title,
                it.publish_time_str,
                it.product_name,
                it.cleaned_name,
                it.price,
                it.video_url,
                it.detail_url,
            ]
        )
    return rows


def _build_account_rows(
    accounts: List[AccountTarget],
    all_items: List[VideoProductItem],
) -> List[List[Any]]:
    rows = []
    accounts = deduplicate_accounts(accounts)
    all_items = deduplicate_items(all_items)
    for acc in accounts:
        identity = stable_id_for(acc)
        unambiguous_name = (
            bool(acc.name)
            and sum(
                other.platform == acc.platform and other.name == acc.name
                for other in accounts
            )
            == 1
        )
        acc_items = [
            it
            for it in all_items
            if it.platform == acc.platform
            and (
                (bool(identity) and it.account_id == identity)
                or (
                    not identity
                    and bool(it.account_id)
                    and it.account_id == acc.url_or_id
                )
                or (
                    not it.account_id
                    and unambiguous_name
                    and it.account_name == acc.name
                )
            )
        ]
        acc_videos = set(it.video_id for it in acc_items)
        rows.append(
            [
                acc.platform,
                acc.name or acc.url_or_id,
                account_profile_url(acc),
                acc.status,
                len(acc_videos),
                len(acc_items),
            ]
        )
    return rows


def _write_sheet(worksheet, headers: List[str], rows: List[List[Any]]) -> None:
    worksheet.append(headers)
    for row_index, row in enumerate(rows, start=2):
        worksheet.append(row)
        for column, value in enumerate(row, start=1):
            if isinstance(value, str) and value.startswith("="):
                worksheet.cell(row=row_index, column=column).data_type = "s"

    # Keep the column names visible and let users filter/sort each report field.
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    header_fill = PatternFill(
        start_color="1F4E79", end_color="1F4E79", fill_type="solid"
    )
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Microsoft YaHei", size=10)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    style_name = "jingxu-body"
    if style_name not in worksheet.parent.named_styles:
        worksheet.parent.add_named_style(
            NamedStyle(
                name=style_name,
                font=cell_font,
                border=thin_border,
                alignment=Alignment(vertical="center"),
            )
        )

    worksheet.views.sheetView[0].showGridLines = True

    for col_idx in range(1, len(headers) + 1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for cell in worksheet[col_letter]:
            if cell.row > 1:
                cell.style = style_name
                if isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
            val_str = str(cell.value or "")
            char_len = sum(2 if ord(c) > 127 else 1 for c in val_str)
            if char_len > max_len:
                max_len = char_len
        worksheet.column_dimensions[col_letter].width = max(10, min(max_len + 4, 60))

    worksheet.row_dimensions[1].height = 26
    for r in range(2, worksheet.max_row + 1):
        worksheet.row_dimensions[r].height = 20


class DataExporter:
    """Exports crawled data and retained product analysis into Excel, CSV, or JSON."""

    @staticmethod
    def export_to_excel(
        file_path: str,
        retained_products: List[RetainedProduct],
        all_items: List[VideoProductItem],
        accounts: List[AccountTarget] = None,
    ) -> str:
        """Export analysis results to a styled Excel file with multiple sheets."""
        all_items = deduplicate_items(all_items)
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)

        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = f"保留相同商品汇总({len(retained_products)}款)"
        _write_sheet(
            summary_sheet, RETAINED_HEADERS, _build_retained_rows(retained_products)
        )

        detail_sheet = workbook.create_sheet(f"全部带货商品明细({len(all_items)}条)")
        _write_sheet(detail_sheet, DETAIL_HEADERS, _build_detail_rows(all_items))

        account_rows = _build_account_rows(accounts or [], all_items)
        if account_rows:
            account_sheet = workbook.create_sheet("对标账号爬取统计")
            _write_sheet(account_sheet, ACCOUNT_HEADERS, account_rows)

        workbook.save(file_path)
        return file_path

    @staticmethod
    def export_to_csv(file_path: str, retained_products: List[RetainedProduct]) -> str:
        """Export retained products to CSV."""
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        with open(file_path, "w", encoding="utf_8_sig", newline="") as f:
            writer = csv.DictWriter(
                f, fieldnames=RETAINED_HEADERS, extrasaction="ignore"
            )
            writer.writeheader()
            for idx, p in enumerate(retained_products, 1):
                writer.writerow(
                    {
                        "序号": idx,
                        "商品名称": p.product_name,
                        "清洗归一名称": p.cleaned_name,
                        "涉及平台": p.platform,
                        "共同带货账号数": p.account_count,
                        "带货视频总数": p.video_count,
                        "示例视频点赞量": p.sample_video_like_count,
                        "带货账号列表": ", ".join(p.accounts),
                        "参考价格区间": p.price_range,
                        "示例视频日期": p.sample_video_date,
                        "示例视频链接": p.sample_video_url,
                    }
                )
        return file_path

    @staticmethod
    def export_to_json(
        file_path: str,
        retained_products: List[RetainedProduct],
        all_items: List[VideoProductItem],
        excel_path: Optional[str] = None,
        diagnostics: Optional[Dict[str, Any]] = None,
        accounts: list[AccountTarget] | None = None,
    ) -> str:
        """Export full data to JSON.

        ``diagnostics`` 携带环境、登录态与商品提取诊断；工作台用它把
        「为什么一个商品都没提取到」直接展示给用户，而不是只报「0 件」。
        """
        all_items = deduplicate_items(all_items)
        unique_accounts = deduplicate_accounts(accounts or [])
        incomplete = [
            acc
            for acc in unique_accounts
            if acc.enabled
            and acc.status not in ("已完成", "完成", "completed", "重复账号，已跳过")
        ]
        message = (
            f"抓取未完整结束：{len(incomplete)} 个账号未完成，已保留 "
            f"{len(retained_products)} 款商品、{len(all_items)} 条带货记录。"
            if incomplete
            else f"抓取完成：保留 {len(retained_products)} 款商品，共 {len(all_items)} 条带货记录。"
        )

        def account_data(acc):
            return {
                **acc.to_dict(),
                "status": acc.status,
                "scraped_videos": acc.scraped_videos,
                "found_products": acc.found_products,
            }

        data = {
            "exported_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "retained_count": len(retained_products),
            "total_items_count": len(all_items),
            "excel_path": excel_path or "",
            "retained_products": [p.to_summary_dict() for p in retained_products],
            "all_items": [it.to_dict() for it in all_items],
            "diagnostics": diagnostics or {},
            "accounts": [account_data(acc) for acc in unique_accounts],
            "account_resolutions": [account_data(acc) for acc in (accounts or [])],
            "crawl_complete": not incomplete,
            "message": message,
        }
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        directory = os.path.dirname(os.path.abspath(file_path))
        fd, temporary = tempfile.mkstemp(
            prefix=".jingxu-result-", suffix=".json.tmp", dir=directory
        )
        try:
            with os.fdopen(fd, "w", encoding="utf-8") as stream:
                json.dump(data, stream, ensure_ascii=False, indent=2)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, file_path)
        finally:
            if os.path.exists(temporary):
                os.unlink(temporary)
        return file_path
